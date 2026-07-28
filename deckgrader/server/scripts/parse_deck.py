#!/usr/bin/env python3
"""Parse a .pptx file into structured JSON for DeckGrader.

Usage: python3 parse_deck.py <path-to-deck.pptx>

Prints deck JSON to stdout on success. On failure, prints a JSON error
object to stderr and exits non-zero. Never prints deck content to stderr.
"""

import json
import sys

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    from pptx.util import Emu
except ImportError:
    json.dump({"error": "python-pptx is not installed"}, sys.stderr)
    sys.exit(3)


def fail(message: str, code: int = 1) -> None:
    json.dump({"error": message}, sys.stderr)
    sys.exit(code)


def shape_kind(shape) -> str:
    st = shape.shape_type
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "image"
    if st == MSO_SHAPE_TYPE.CHART or shape.has_chart:
        return "chart"
    if st == MSO_SHAPE_TYPE.TABLE or shape.has_table:
        return "table"
    if shape.has_text_frame:
        return "textBox"
    return "other"


def collect_text_frames(shapes):
    """Yield (shape, text_frame) pairs, descending into groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from collect_text_frames(shape.shapes)
        elif shape.has_text_frame:
            yield shape, shape.text_frame


def collect_shapes(shapes):
    """Yield leaf shapes, descending into groups."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from collect_shapes(shape.shapes)
        else:
            yield shape


def parse_slide(slide, index: int) -> dict:
    title = None
    has_title_placeholder = False
    title_shape_id = None
    try:
        title_shape = slide.shapes.title
    except Exception:
        title_shape = None
    if title_shape is not None:
        has_title_placeholder = True
        title_shape_id = title_shape.shape_id
        text = title_shape.text_frame.text.strip()
        title = text if text else None

    body_text: list[str] = []
    font_sizes: set[float] = set()
    fonts: set[str] = set()
    bullet_depth_max = 0
    counts = {"textBoxes": 0, "images": 0, "charts": 0, "tables": 0}

    for shape in collect_shapes(slide.shapes):
        kind = shape_kind(shape)
        if kind == "image":
            counts["images"] += 1
        elif kind == "chart":
            counts["charts"] += 1
        elif kind == "table":
            counts["tables"] += 1
            for row in shape.table.rows:
                for cell in row.cells:
                    cell_text = cell.text.strip()
                    if cell_text:
                        body_text.append(cell_text)
        elif kind == "textBox":
            counts["textBoxes"] += 1

        if not getattr(shape, "has_text_frame", False):
            continue
        is_title = shape.shape_id == title_shape_id
        for para in shape.text_frame.paragraphs:
            para_text = "".join(run.text for run in para.runs).strip()
            if para_text and not is_title:
                body_text.append(para_text)
                bullet_depth_max = max(bullet_depth_max, para.level + 1)
            for run in para.runs:
                if run.font.size is not None:
                    font_sizes.add(round(Emu(run.font.size).pt, 1))
                if run.font.name:
                    fonts.add(run.font.name)

    notes = None
    if slide.has_notes_slide:
        notes_text = slide.notes_slide.notes_text_frame.text.strip()
        notes = notes_text if notes_text else None

    title_words = len(title.split()) if title else 0
    word_count = title_words + sum(len(t.split()) for t in body_text)

    return {
        "index": index,
        "title": title,
        "bodyText": body_text,
        "notes": notes,
        "shapeCounts": counts,
        "wordCount": word_count,
        "bulletDepthMax": bullet_depth_max,
        "hasTitlePlaceholder": has_title_placeholder,
        "fontSizesUsed": sorted(font_sizes),
        "distinctFonts": sorted(fonts),
    }


def main() -> None:
    if len(sys.argv) != 2:
        fail("usage: parse_deck.py <path-to-deck.pptx>", 2)

    path = sys.argv[1]
    try:
        prs = Presentation(path)
    except FileNotFoundError:
        fail("file not found", 2)
    except Exception:
        fail("could not open file as a PowerPoint presentation (corrupt or wrong format)", 1)

    try:
        slides = [parse_slide(slide, i + 1) for i, slide in enumerate(prs.slides)]
    except Exception:
        fail("failed while reading slide contents (deck may be malformed)", 1)

    all_fonts = sorted({f for s in slides for f in s["distinctFonts"]})
    total_words = sum(s["wordCount"] for s in slides)
    deck = {
        "slideCount": len(slides),
        "slides": slides,
        "deckStats": {
            "totalWordCount": total_words,
            "avgWordsPerSlide": round(total_words / len(slides)) if slides else 0,
            "distinctFontsCount": len(all_fonts),
            "slidesMissingTitles": sum(1 for s in slides if not s["title"]),
        },
    }
    json.dump(deck, sys.stdout)


if __name__ == "__main__":
    main()
