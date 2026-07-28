#!/usr/bin/env python3
"""Render a DeckGrader slide spec into a .pptx.

Usage: python3 build_deck.py <output-path.pptx> < spec.json

Spec (stdin JSON):
{
  "deckTitle": "string",
  "subtitle": "string | null",
  "slides": [
    {
      "title": "action-title string",
      "bullets": [{"text": "string", "level": 0}],
      "notes": "string | null"
    }
  ]
}

Follows the rubric's own craft rules: action title in the title placeholder,
bullet depth capped at 2 levels, single font family (template default).
On failure prints a JSON error to stderr and exits non-zero. Never prints
deck content to stderr.
"""

import json
import sys

try:
    from pptx import Presentation
    from pptx.util import Pt
except ImportError:
    json.dump({"error": "python-pptx is not installed"}, sys.stderr)
    sys.exit(3)


def fail(message: str, code: int = 1) -> None:
    json.dump({"error": message}, sys.stderr)
    sys.exit(code)


def add_content_slide(prs, title: str, bullets, notes) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[1])  # title + content
    slide.shapes.title.text = title
    slide.shapes.title.text_frame.word_wrap = True
    for para in slide.shapes.title.text_frame.paragraphs:
        para.font.size = Pt(24)

    body = slide.placeholders[1].text_frame
    body.clear()
    body.word_wrap = True
    first = True
    for bullet in bullets:
        text = str(bullet.get("text", "")).strip()
        if not text:
            continue
        para = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        para.text = text
        para.level = max(0, min(1, int(bullet.get("level", 0))))  # depth <= 2 levels
        para.font.size = Pt(16 if para.level == 0 else 14)

    if notes:
        slide.notes_slide.notes_text_frame.text = str(notes)


def main() -> None:
    if len(sys.argv) != 2:
        fail("usage: build_deck.py <output-path.pptx> < spec.json", 2)

    try:
        spec = json.load(sys.stdin)
    except json.JSONDecodeError:
        fail("spec on stdin is not valid JSON", 2)

    deck_title = str(spec.get("deckTitle") or "").strip()
    slides = spec.get("slides")
    if not deck_title or not isinstance(slides, list) or not slides:
        fail("spec must include deckTitle and a non-empty slides array", 2)

    prs = Presentation()

    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_slide.shapes.title.text = deck_title
    subtitle = spec.get("subtitle")
    if subtitle:
        title_slide.placeholders[1].text = str(subtitle)

    try:
        for slide in slides:
            add_content_slide(
                prs,
                str(slide.get("title") or "Untitled"),
                slide.get("bullets") or [],
                slide.get("notes"),
            )
        prs.save(sys.argv[1])
    except Exception:
        fail("failed to render the deck", 1)


if __name__ == "__main__":
    main()
