#!/usr/bin/env python3
"""Generate test fixture decks for DeckGrader.

Creates two decks in server/fixtures/:
  bad_deck.pptx  — label titles, walls of text, deep bullets, no recommendation
  good_deck.pptx — action titles forming a coherent argument, ends with next steps
"""

import os

from pptx import Presentation
from pptx.util import Inches, Pt

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), "..", "fixtures")

LOREM = (
    "The market environment continues to evolve in a number of directions and "
    "there are many factors that could be considered relevant when thinking "
    "about the overall strategic positioning of the business going forward. "
)


def add_slide(prs, title_text, bullets, deep=False):
    layout = prs.slide_layouts[1]  # title + content
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = title_text
    body = slide.placeholders[1].text_frame
    body.clear()
    first = True
    for level, text in bullets:
        para = body.paragraphs[0] if first else body.add_paragraph()
        first = False
        para.text = text
        para.level = min(level, 4) if deep else min(level, 1)
        para.font.size = Pt(14)
    return slide


def make_bad_deck(path):
    prs = Presentation()

    # Title slide
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "Project Update"
    slide.placeholders[1].text = "Q3 Materials"

    wall = (LOREM * 5).split(". ")
    add_slide(prs, "Overview", [(0, s + ".") for s in wall[:8] if s], deep=True)
    add_slide(
        prs,
        "Background",
        [
            (0, "There is a lot of context that matters here"),
            (1, "Many stakeholders have many opinions"),
            (2, "Some of these opinions are quite significant"),
            (3, "A large number of meetings were held"),
            (4, "Numerous action items were generated in those meetings"),
            (0, LOREM),
            (0, LOREM),
        ],
        deep=True,
    )
    add_slide(prs, "Data", [(0, s + ".") for s in (LOREM * 5).split(". ")[:9] if s], deep=True)
    add_slide(
        prs,
        "Analysis",
        [
            (0, "We looked at many things and found various results"),
            (1, "Some segments performed well, others less so"),
            (2, "Performance was significant in certain areas"),
            (3, "Many customers said many different things"),
            (0, LOREM),
            (0, LOREM),
            (0, LOREM),
        ],
        deep=True,
    )
    add_slide(prs, "Findings", [(0, s + ".") for s in (LOREM * 5).split(". ")[:9] if s], deep=True)
    add_slide(prs, "Thank You", [(0, "Questions?")])

    prs.save(path)


def make_good_deck(path):
    prs = Presentation()

    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "SMB churn is fixable: a 90-day retention plan for the CEO"
    slide.placeholders[1].text = "Prepared for the executive committee — decision requested on option A"

    add_slide(
        prs,
        "Churn rose to 4.1% monthly in Q2, driven almost entirely by SMB accounts",
        [
            (0, "Monthly churn: 2.3% (Q4) to 3.1% (Q1) to 4.1% (Q2)"),
            (0, "SMB segment accounts for 86% of the increase; enterprise churn flat at 0.8%"),
            (0, "Source: billing system cohort export, June 30"),
        ],
    )
    add_slide(
        prs,
        "Churn concentrates in month-2 SMB cohorts that never completed onboarding",
        [
            (0, "61% of churned SMB accounts never finished the 3-step onboarding flow"),
            (0, "Accounts completing onboarding churn at 1.2% vs 6.8% for those who do not"),
            (0, "Month 2 is the cliff: 74% of SMB churn occurs in days 31-60"),
        ],
    )
    add_slide(
        prs,
        "A guided onboarding intervention would cut SMB churn roughly in half",
        [
            (0, "Pilot with 120 accounts: completion rose from 39% to 82%"),
            (0, "Pilot cohort churn: 2.9% vs 6.1% control over the same 60 days"),
            (0, "Cost: 2 CS hires (~$180K/yr) vs ~$1.4M annual revenue retained"),
        ],
    )
    add_slide(
        prs,
        "Recommendation: fund the onboarding team now; decision needed by August 15",
        [
            (0, "Approve 2 CS onboarding hires and the in-app guide build (option A)"),
            (0, "Next steps: hiring req out this week; instrument cohort dashboard by Aug 1"),
            (0, "Review results at the October board meeting"),
        ],
    )

    prs.save(path)


def main():
    os.makedirs(FIXTURES_DIR, exist_ok=True)
    bad = os.path.join(FIXTURES_DIR, "bad_deck.pptx")
    good = os.path.join(FIXTURES_DIR, "good_deck.pptx")
    make_bad_deck(bad)
    make_good_deck(good)
    print(f"wrote {bad}")
    print(f"wrote {good}")


if __name__ == "__main__":
    main()
